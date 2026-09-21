import os
import re
import unicodedata
from datetime import datetime
import openpyxl

TEMPLATE_PATH = "MATRICE FRAIS.xlsx"
SHEET_NAME = "Dépenses Pavillons"

def normalize_str(s: str) -> str:
    if not s:
        return ""
    s = s.lower().strip()
    s = unicodedata.normalize('NFD', s).encode('ascii', 'ignore').decode("utf-8")
    return re.sub(r'[^a-z0-9]', '', s)

def parse_montant(montant_str) -> float:
    if isinstance(montant_str, (int, float)):
        return float(montant_str)
    if not montant_str:
        return 0.0
    clean = str(montant_str).replace(',', '.').replace(' ', '')
    match = re.search(r'[-+]?\d*\.?\d+', clean)
    return float(match.group()) if match else 0.0

def parse_date(date_val):
    if not date_val:
        return None
    if isinstance(date_val, datetime):
        return date_val
    try:
        # Format ISO (ex: YYYY-MM-DD depuis Supabase)
        return datetime.strptime(str(date_val)[:10], "%Y-%m-%d")
    except ValueError:
        try:
            # Format classique DD/MM/YYYY
            return datetime.strptime(str(date_val).strip(), "%d/%m/%Y")
        except ValueError:
            return str(date_val).strip()

def get_category_column_map(sheet) -> dict:
    category_map = {}
    for col in range(7, 25):  # Colonnes G (7) à X (24)
        val = sheet.cell(row=7, column=col).value
        if val:
            category_map[normalize_str(str(val))] = col
    return category_map

def find_best_col(extracted_cat: str, category_map: dict) -> int | None:
    if not extracted_cat:
        return None
    cat_norm = normalize_str(extracted_cat)

    if cat_norm in category_map:
        return category_map[cat_norm]

    for key, col in category_map.items():
        if cat_norm in key or key in cat_norm:
            return col

    words_extracted = set(re.findall(r'\w+', cat_norm))
    best_score, best_col = 0, None
    for key, col in category_map.items():
        words_key = set(re.findall(r'\w+', key))
        common = sum(1 for w1 in words_extracted for w2 in words_key if w1[:4] == w2[:4])
        if common > best_score:
            best_score = common
            best_col = col

    return best_col

def generate_monthly_excel(all_records: list[dict], month_year_str: str) -> str:
    """
    Prend toutes les dépenses d'un mois données par Supabase,
    les injecte dans une copie propre du template vierge,
    et sauvegarde sous 'MATRICE_FRAIS_MM_YYYY.xlsx'.
    """
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(f"Le fichier modèle '{TEMPLATE_PATH}' est introuvable à la racine.")

    # Toujours partir du modèle vierge
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    sheet = wb[SHEET_NAME]
    cat_map = get_category_column_map(sheet)

    row = 9  # Première ligne de saisie du tableau

    for fiche in all_records:
        if row > 35:
            raise ValueError("Le tableau Excel est plein (limite de 27 lignes/mois atteinte).")

        montant_val = parse_montant(fiche.get("montant"))

        # Colonne A : Date
        sheet.cell(row=row, column=1).value = parse_date(fiche.get("date"))

        # Colonne B : Facture / Libellé
        sheet.cell(row=row, column=2).value = fiche.get("facture_nom")

        # Colonne C : N° de page
        num_page = fiche.get("numero_page")
        try:
            sheet.cell(row=row, column=3).value = int(num_page) if num_page is not None else None
        except (ValueError, TypeError):
            sheet.cell(row=row, column=3).value = num_page

        # Colonne D : Montant Global
        sheet.cell(row=row, column=4).value = montant_val

        # Colonnes G à X : Ventillation par catégorie
        target_col = find_best_col(fiche.get("categorie"), cat_map)
        if target_col:
            sheet.cell(row=row, column=target_col).value = montant_val

        row += 1

    output_filename = f"MATRICE_FRAIS_{month_year_str}.xlsx"
    wb.save(output_filename)
    return output_filename